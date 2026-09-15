from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = r"C:\Users\wang0\Documents\Codex\2026-07-29\santo0520-ncku-llm-drug-discovery-assistant\outputs\NCKU_Open_Model_Drug_Discovery_Assistant\Agent1_操作手冊.docx"
PY = r'D:\NCKU_Drug_Discovery\.venv\Scripts\python.exe'
AGENT = r'C:\Users\wang0\Documents\Codex\2026-07-29\santo0520-ncku-llm-drug-discovery-assistant\outputs\NCKU_Open_Model_Drug_Discovery_Assistant\agent.py'
OUTDIR = r'D:\NCKU_Drug_Discovery\outputs'
PLANDIR = r'D:\NCKU_Drug_Discovery\outputs\plans'
OLLAMA = r'C:\Users\wang0\Documents\Codex\2026-07-29\santo0520-ncku-llm-drug-discovery-assistant\work\ollama-portable\ollama.exe'
BLUE, NAVY, LIGHT, GRAY, RED = "2E74B5", "203A5F", "E8EEF5", "667085", "B42318"


def shade(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def widths(table, values):
    table.autofit = False
    for row in table.rows:
        for cell, value in zip(row.cells, values):
            cell.width = Inches(value)


def code(doc, text):
    p = doc.add_paragraph(style="Code Block")
    p.paragraph_format.keep_together = True
    for line in text.splitlines():
        r = p.add_run(line + "\n")
        r.font.name = "Consolas"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
        r.font.size = Pt(8.5)


def bullet(doc, text):
    doc.add_paragraph(text, style="List Bullet")


def note(doc, title, text, warning=False):
    table = doc.add_table(rows=1, cols=1)
    widths(table, [6.5])
    shade(table.cell(0, 0), "FDECEC" if warning else "EDF4FB")
    p = table.cell(0, 0).paragraphs[0]
    r = p.add_run(title + "　")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(RED if warning else BLUE)
    p.add_run(text)
    doc.add_paragraph()


def matrix(doc, headers, rows, col_widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        table.cell(0, i).text = header
        shade(table.cell(0, i), LIGHT)
        table.cell(0, i).paragraphs[0].runs[0].bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    widths(table, col_widths)
    return table


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(0.8), Inches(0.75)
sec.left_margin = sec.right_margin = Inches(1)
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.2
for name, size, color in (("Heading 1", 16, BLUE), ("Heading 2", 13, BLUE), ("Heading 3", 11.5, NAVY)):
    style = doc.styles[name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.keep_with_next = True
code_style = doc.styles.add_style("Code Block", 1)
code_style.font.name, code_style.font.size = "Consolas", Pt(8.5)
code_style.paragraph_format.left_indent = Inches(0.18)
code_style.paragraph_format.right_indent = Inches(0.12)
code_style.paragraph_format.space_after = Pt(8)
shd = OxmlElement("w:shd")
shd.set(qn("w:fill"), "F3F5F7")
code_style.element.get_or_add_pPr().append(shd)
header = sec.header.paragraphs[0]
header.text = "NCKU Open Model Drug Discovery Assistant｜Agent 1"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header.runs[0].font.size = Pt(8)
header.runs[0].font.color.rgb = RGBColor.from_string(GRAY)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run("研究開發操作手冊｜2026-08-14")
r.font.size = Pt(8)
r.font.color.rgb = RGBColor.from_string(GRAY)

# Editorial cover
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(100)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("AGENT 1")
r.bold, r.font.size = True, Pt(13)
r.font.color.rgb = RGBColor.from_string(BLUE)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("操作手冊")
r.bold, r.font.size = True, Pt(30)
r.font.color.rgb = RGBColor.from_string(NAVY)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("確定性資料庫查詢與 Qwen Research Planner")
r.font.size = Pt(15)
r.font.color.rgb = RGBColor.from_string(GRAY)
p.paragraph_format.space_after = Pt(42)
note(doc, "目前版本重點", "保留不經 LLM 的可重現查詢；另新增由 Qwen 規劃、程式驗證的研究流程。兩條路徑互不取代。")
p = doc.add_paragraph("正式程式目錄")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(65)
p.runs[0].bold = True
p.runs[0].font.color.rgb = RGBColor.from_string(BLUE)
p = doc.add_paragraph(AGENT.rsplit('\\', 1)[0])
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.runs[0].font.size = Pt(8)
p.runs[0].font.color.rgb = RGBColor.from_string(GRAY)

doc.add_page_break()
doc.add_heading("1. 先選擇正確模式", level=1)
matrix(doc, ["指令", "LLM", "用途", "目前狀態"], [
    ("discover-targets", "否", "疾病 → 前 N 個 Targets", "可執行"),
    ("agent1", "否", "指定 Disease–Target 原始證據", "可執行"),
    ("plan", "是", "只建立並驗證 Research Plan", "不收集正式證據"),
    ("agent1-planned", "是", "進入受控 Agent 1 流程", "停在 collection policy gate"),
], [1.25, .65, 2.55, 2.05])
doc.add_heading("使用原則", level=2)
bullet(doc, "和專家討論原始欄位、資料量或儲存結構：使用不經 LLM 的指令。")
bullet(doc, "檢查 Qwen 如何安排四個來源：使用 plan。")
bullet(doc, "測試完整自主流程框架：使用 agent1-planned；目前不會無限制下載。")
note(doc, "重要", "Open Targets 分數僅作初始排序參考，不是因果、療效、臨床成功率或專案信心分數。")

doc.add_heading("2. 執行前準備", level=1)
doc.add_paragraph("指令可在任何 PowerShell 目錄執行，因為以下命令使用完整路徑。Python 環境與輸出均放在 D 槽。")
doc.add_heading("只有 LLM 模式需要啟動 Ollama", level=2)
code(doc, f'$env:OLLAMA_MODELS="D:\\OllamaModels"\n& "{OLLAMA}" serve')
doc.add_paragraph("保持此 PowerShell 視窗開啟，再開第二個 PowerShell 執行 plan 或 agent1-planned。")
code(doc, f'& "{OLLAMA}" --version\n& "{OLLAMA}" list')

doc.add_page_break()
doc.add_heading("3. 不經 LLM：疾病探索", level=1)
doc.add_paragraph("用途：解析疾病 ID，依 Open Targets overall association score 取得前 10 個 Targets。")
code(doc, f'& "{PY}" "{AGENT}" discover-targets `\n  --disease "pancreatic cancer" `\n  --top 10 `\n  --output "{OUTDIR}"')
doc.add_paragraph("預期檔案：")
code(doc, OUTDIR + r"\pancreatic_cancer_top10_targets.json")
bullet(doc, "檔名依 --disease 自動產生。")
bullet(doc, "不需要 Ollama；資料由 Open Targets 工具取得。")

doc.add_heading("4. 不經 LLM：指定 Disease–Target", level=1)
code(doc, f'& "{PY}" "{AGENT}" agent1 `\n  --disease "Alzheimer\'s disease" `\n  --target "APP" `\n  --max-evidence-records 500 `\n  --output "{OUTDIR}\\agent1_APP_expert_review.json"')
bullet(doc, "不加 --with-synthesizer：完全不使用 LLM。")
bullet(doc, "加 --with-synthesizer：只有整理階段使用 Qwen；原始資料仍由工具取得。")

doc.add_page_break()
doc.add_heading("5. 使用 Qwen：只產生 Research Plan", level=1)
doc.add_paragraph("plan 先用確定性工具解析 ID；疾病模式還會先取得前 10 個 Targets。之後才交給 Qwen 規劃，不執行正式證據收集。")
doc.add_heading("只輸入疾病", level=2)
code(doc, f'& "{PY}" "{AGENT}" plan `\n  --disease "pancreatic cancer" `\n  --top 10 `\n  --output "{PLANDIR}"')
doc.add_heading("疾病＋指定 Target", level=2)
code(doc, f'& "{PY}" "{AGENT}" plan `\n  --disease "pancreatic cancer" `\n  --target "KRAS" `\n  --output "{PLANDIR}"')
doc.add_paragraph("預期檔案：")
code(doc, PLANDIR + r"\pancreatic_cancer_research_plan.json")
doc.add_heading("計畫驗證規則", level=2)
for text in [
    "每個 Target 必須包含 Open Targets、Europe PMC、ClinicalTrials.gov、ChEMBL。",
    "只能使用已註冊工具，不能取消必要來源。",
    "疾病與 Target ID 由程式驗證，Qwen 不得修改。",
    "Qwen 新增搜尋詞必須標示為 unvalidated_search_term。",
    "格式錯誤最多交回 Qwen 修正 2 次；仍失敗就停止。",
]:
    bullet(doc, text)

doc.add_heading("6. 使用 Qwen：受控 Agent 1 流程", level=1)
code(doc, f'& "{PY}" "{AGENT}" agent1-planned `\n  --disease "pancreatic cancer" `\n  --top 10 `\n  --output "{PLANDIR}"')
note(doc, "目前限制", "各來源 Collection Policy 尚未經資料量實測確定。因此會驗證並保存計畫，但以 collection_policy_not_configured 停止，不會大量下載。", True)

doc.add_page_break()
doc.add_heading("7. Agent 1 內部責任", level=1)
matrix(doc, ["角色", "方式", "責任"], [
    ("Research Planner", "Qwen", "規劃關鍵字、參數、順序、問題與理由"),
    ("Tool Runner", "Python", "依合法計畫查詢，不得編造結果"),
    ("Evidence Validator", "Python", "檢查 ID、來源、格式、錯誤、分頁與截斷"),
    ("Evidence Synthesizer", "Qwen", "判斷重要性、支持／衝突及資料缺口"),
], [1.55, 1.05, 3.9])
doc.add_heading("兩輪查詢原則", level=2)
for text in [
    "第一輪：每個 Target 的四個必要來源全部查詢。",
    "Evidence Synthesizer 檢查資料缺口。",
    "必要時只允許一輪追加查詢，避免無限循環。",
    "Collection Policy 之後依各來源實測筆數、大小與時間決定。",
]:
    bullet(doc, text)

doc.add_heading("8. 常見錯誤", level=1)
matrix(doc, ["訊息／現象", "處理方式"], [
    ("ollama 無法辨識", "使用 ollama.exe 完整路徑，不依賴 PATH。"),
    ("Unable to create process", "D 槽 venv 綁定的基礎 Python 已移除；需修復或重建。"),
    ("Open Targets request failed", "確認網路、API 狀態與錯誤內容。"),
    ("research_planner_failed", "確認 Ollama 已啟動、Qwen 模型存在。"),
    ("collection_policy_not_configured", "預期的安全停止；等待各來源上限完成實測。"),
], [2.2, 4.3])

doc.add_heading("9. 建議操作順序", level=1)
for text in [
    "使用 discover-targets 取得疾病候選 Targets。",
    "使用 agent1 取得可重現的指定 Disease–Target 原始資料。",
    "將 expert review JSON 提供給領域專家，確認分類與比較方式。",
    "使用 plan 檢查 Qwen 的查詢規劃品質。",
    "量測四個來源的筆數、檔案大小與時間，再制定 Collection Policy。",
    "確認 Collection Policy 後，才開放 agent1-planned 完整收集。",
]:
    doc.add_paragraph(text, style="List Number")

doc.save(OUT)
print(OUT)
